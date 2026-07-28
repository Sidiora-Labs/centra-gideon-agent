"""DFE-2: run-level .docx fidelity, asserted on the REOPENED document.

Every claim here is a round trip: render to bytes, hand the bytes back to python-docx,
and read the property off the document. Asserting the model instead would prove only
that we called `add_run`, which is exactly the failure this atom exists to rule out.

`Run` / `ParagraphStyle` / `Cell` / `PageSetup` are declared by `documents/model.py`; the
writer only ever READS attributes off them, so the stand-ins below stand in for the real
dataclasses and let this suite pin the writer's contract on its own.
"""

from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from gideon.documents.model import BLOCK_KINDS, Block, DocumentModel
from gideon.documents.writers import docx_writer
from gideon.documents.writers.docx_writer import render_docx
from gideon.knowledge.readers import FileReader

_MONOSPACE = "Courier New"
_URL = "https://example.invalid/docs"


@dataclass
class _Run:
    text: str = ""
    bold: bool = False
    italic: bool = False
    code: bool = False
    link: str = ""


@dataclass
class _Style:
    align: str = ""
    space_before_pt: float = 0.0
    space_after_pt: float = 0.0
    line_spacing: float = 0.0


@dataclass
class _Cell:
    runs: list = field(default_factory=list)
    text: str = ""
    bold: bool = False
    align: str = ""


@dataclass
class _Page:
    orientation: str = ""
    margin_in: float = 0.0


def _block(kind: str, *, runs=(), cells=(), style=None, **kw) -> Block:
    """A `Block` carrying the inline fields the writer reads."""
    block = Block(kind=kind, **kw)
    block.runs = list(runs)
    block.cells = [list(row) for row in cells]
    block.style = style
    return block


def _model(*blocks: Block, title: str = "", page=None) -> DocumentModel:
    model = DocumentModel(title=title, blocks=list(blocks))
    model.page = page
    return model


def _reopen(data: bytes):
    """The bytes, read back by python-docx — the only surface these tests assert on."""
    return Document(io.BytesIO(data))


def _document_xml(data: bytes) -> bytes:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return archive.read("word/document.xml")


def _part_names(data: bytes) -> list[str]:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return sorted(archive.namelist())


# --------------------------------------------------------------------------- the clause


def test_a_bold_run_reads_back_bold_via_python_docx():
    """DFE-2's clause, round-tripped through the file rather than the model."""
    data = render_docx(
        _model(_block("paragraph", runs=[_Run("plain "), _Run("strong", bold=True)]))
    )

    para = _reopen(data).paragraphs[0]
    assert [run.text for run in para.runs] == ["plain ", "strong"]
    assert para.runs[1].bold is True, "a bold run must read back bold from the document"
    assert para.runs[0].bold is None, "a plain run inherits; it must not declare bold OFF"


def test_an_italic_run_reads_back_italic_via_python_docx():
    data = render_docx(
        _model(_block("paragraph", runs=[_Run("plain "), _Run("leaning", italic=True)]))
    )

    para = _reopen(data).paragraphs[0]
    assert para.runs[1].italic is True
    assert para.runs[0].italic is None


def test_a_bold_run_in_a_heading_reads_back_bold():
    data = render_docx(_model(_block("heading", level=2, runs=[_Run("Big", bold=True)])))

    para = _reopen(data).paragraphs[0]
    assert para.style.name == "Heading 2"
    assert para.runs[0].text == "Big"
    assert para.runs[0].bold is True


def test_a_code_run_reads_back_in_a_monospace_font():
    runs = [_Run("call "), _Run("render_docx", code=True)]
    data = render_docx(_model(_block("paragraph", runs=runs)))

    para = _reopen(data).paragraphs[0]
    assert para.runs[1].font.name == _MONOSPACE
    assert para.runs[0].font.name is None, "only the code run is respelled"


def test_a_link_run_keeps_its_display_text_and_its_url():
    """A dropped URL is lost user content, so both halves are asserted."""
    runs = [_Run("see "), _Run("the docs", link=_URL)]
    data = render_docx(_model(_block("paragraph", runs=runs)))

    para = _reopen(data).paragraphs[0]
    # A hyperlink's run is a child of w:hyperlink, not of w:p, so it is deliberately
    # absent from para.runs — the visible text arrives through para.text.
    assert para.text == "see the docs"
    assert [link.address for link in para.hyperlinks] == [_URL]
    assert [run.text for run in para.hyperlinks[0].runs] == ["the docs"]


def test_a_link_run_with_no_display_text_shows_its_url():
    data = render_docx(_model(_block("paragraph", runs=[_Run(link=_URL)])))

    para = _reopen(data).paragraphs[0]
    assert para.text == _URL, "an empty clickable span would be invisible"
    assert [link.address for link in para.hyperlinks] == [_URL]


def test_the_hyperlink_relationship_is_external_and_points_at_the_url():
    data = render_docx(_model(_block("paragraph", runs=[_Run("the docs", link=_URL)])))

    doc = _reopen(data)
    rel = next(r for r in doc.part.rels.values() if r.reltype.endswith("/hyperlink"))
    assert rel.is_external is True, "an internal relationship would resolve to nothing"
    assert rel.target_ref == _URL


def test_the_hyperlink_run_properties_are_in_ooxml_schema_order():
    """`CT_RPr` is a sequence, not a choice — out-of-order children read as corrupt."""
    runs = [_Run("hot", bold=True, italic=True, code=True, link=_URL)]
    xml = _document_xml(render_docx(_model(_block("paragraph", runs=runs)))).decode("utf-8")

    props = xml.split("<w:rPr>")[1].split("</w:rPr>")[0]
    order = [chunk.split(" ")[0].split("/")[0] for chunk in props.split("<w:")[1:]]
    assert order == ["rFonts", "b", "i", "color", "u"]


def test_a_bold_link_run_reads_back_bold():
    data = render_docx(_model(_block("paragraph", runs=[_Run("hot", bold=True, link=_URL)])))

    para = _reopen(data).paragraphs[0]
    assert para.hyperlinks[0].runs[0].bold is True


# ------------------------------------------------------------------- backwards fidelity


def _legacy_docx(model: DocumentModel) -> bytes:
    """The pre-DFE-2 render, transcribed call-for-call from the writer this atom changes.

    Not a golden file: the SAME model driven through the SAME python-docx calls, so any
    element the runs path adds to a runs-less block turns up as an XML diff.
    """
    doc = Document()
    if model.title:
        doc.add_heading(model.title, level=0)
    for block in model.blocks:
        kind = block.kind
        if kind == "heading":
            doc.add_heading(block.text, level=block.level)
        elif kind == "paragraph":
            doc.add_paragraph(block.text)
        elif kind == "bullets":
            for item in block.items:
                doc.add_paragraph(item, style="List Bullet")
        elif kind == "numbered":
            for item in block.items:
                doc.add_paragraph(item, style="List Number")
        elif kind == "table":
            rows = block.rows
            if rows:
                width = max(len(row) for row in rows)
                table = doc.add_table(rows=len(rows), cols=width)
                table.style = "Table Grid"
                for r, row in enumerate(rows):
                    for c in range(width):
                        table.cell(r, c).text = str(row[c]) if c < len(row) else ""
                for cell in table.rows[0].cells:
                    for para in cell.paragraphs:
                        for run in para.runs:
                            run.bold = True
        elif kind == "code":
            doc.add_paragraph().add_run(block.text).font.name = _MONOSPACE
        elif kind == "pagebreak":
            doc.add_page_break()
        elif kind == "image":
            doc.add_paragraph(f"[image: {block.artifact_slug}]")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _every_kind_runs_less() -> list[Block]:
    return [
        _block("heading", text="Title Two", level=2),
        _block("paragraph", text="Prose."),
        _block("bullets", items=["one", "two"]),
        _block("numbered", items=["first", "second"]),
        _block("table", rows=[["h1", "h2"], ["a"], ["b", "c"]]),
        _block("code", text="x = 1\n"),
        _block("pagebreak"),
        _block("image", artifact_slug="chart-1"),
    ]


def test_a_runs_less_model_renders_exactly_as_it_did_before():
    """Backwards compatibility, proved on the OUTPUT rather than argued about."""
    model = _model(*_every_kind_runs_less(), title="Doc")

    assert _document_xml(render_docx(model)) == _document_xml(_legacy_docx(model))
    assert _part_names(render_docx(model)) == _part_names(_legacy_docx(model))


def test_a_runs_less_paragraph_carries_no_run_properties():
    """The narrow version of the claim above: nothing is added, not even an empty rPr."""
    para = _reopen(render_docx(_model(_block("paragraph", text="Prose.")))).paragraphs[0]

    assert para.runs[0].bold is None
    assert para.runs[0].italic is None
    assert para.runs[0].font.name is None
    assert b"w:rPr" not in _document_xml(render_docx(_model(_block("paragraph", text="Prose."))))


# ----------------------------------------------------------- 0.0 means "writer default"


def _paragraph_with_style(style) -> object:
    return _reopen(render_docx(_model(_block("paragraph", text="a", style=style)))).paragraphs[0]


def test_zero_spacing_means_writer_default_and_never_an_explicit_zero():
    unset = _paragraph_with_style(_Style())
    tight = _paragraph_with_style(_Style(space_before_pt=2, space_after_pt=3, line_spacing=0.9))

    assert unset.paragraph_format.space_before is None, "0.0 must leave the template alone"
    assert unset.paragraph_format.space_after is None
    assert unset.paragraph_format.line_spacing is None
    assert unset.alignment is None

    assert tight.paragraph_format.space_before == Pt(2)
    assert tight.paragraph_format.space_after == Pt(3)
    assert tight.paragraph_format.line_spacing == 0.9
    # "unset" and "tight" must be distinguishable in the OUTPUT, not just in the model.
    assert tight.paragraph_format.space_before != unset.paragraph_format.space_before
    assert tight.paragraph_format.line_spacing != unset.paragraph_format.line_spacing


def test_a_declared_alignment_reads_back_and_an_unknown_one_does_not_guess():
    assert _paragraph_with_style(_Style(align="center")).alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert _paragraph_with_style(_Style(align="right")).alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert _paragraph_with_style(_Style(align="sideways")).alignment is None


def test_a_style_of_none_touches_nothing():
    assert _paragraph_with_style(None).paragraph_format.space_before is None
    assert b"w:pPr" not in _document_xml(render_docx(_model(_block("paragraph", text="a"))))


def _section(page) -> object:
    return _reopen(render_docx(_model(_block("paragraph", text="a"), page=page))).sections[0]


def test_zero_margin_means_writer_default_and_never_a_zero_margin():
    default = _section(None)
    declared = _section(_Page(margin_in=0.75))

    assert _section(_Page()).left_margin == default.left_margin
    assert declared.left_margin == Inches(0.75)
    assert declared.top_margin == Inches(0.75)
    assert declared.left_margin != default.left_margin


def test_landscape_swaps_the_page_and_portrait_is_a_no_op_on_a_portrait_template():
    default = _section(None)
    landscape = _section(_Page(orientation="landscape"))

    assert landscape.page_width == default.page_height
    assert landscape.page_height == default.page_width
    assert _section(_Page(orientation="portrait")).page_width == default.page_width


# ----------------------------------------------------------------------- richer  tables


def test_a_cells_table_keeps_the_row_zero_header_bold_and_honours_per_cell_bold():
    cells = [
        [_Cell(text="h1"), _Cell(text="h2")],
        [_Cell(text="plain"), _Cell(text="loud", bold=True)],
    ]
    table = _reopen(render_docx(_model(_block("table", cells=cells)))).tables[0]

    assert [c.text for c in table.rows[0].cells] == ["h1", "h2"]
    assert all(run.bold for run in table.cell(0, 0).paragraphs[0].runs), "row 0 is the header"
    assert all(run.bold for run in table.cell(0, 1).paragraphs[0].runs)
    assert table.cell(1, 0).paragraphs[0].runs[0].bold is None
    assert table.cell(1, 1).paragraphs[0].runs[0].bold is True


def test_a_cells_table_renders_run_level_formatting_and_alignment():
    cells = [
        [_Cell(runs=[_Run("H", bold=True)])],
        [_Cell(runs=[_Run("it", italic=True)], align="center")],
    ]
    table = _reopen(render_docx(_model(_block("table", cells=cells)))).tables[0]

    assert table.cell(1, 0).paragraphs[0].runs[0].italic is True
    assert table.cell(1, 0).paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER


def test_a_ragged_cells_table_is_normalized_to_its_widest_row():
    cells = [[_Cell(text="a"), _Cell(text="b")], [_Cell(text="c")]]
    table = _reopen(render_docx(_model(_block("table", cells=cells)))).tables[0]

    assert len(table.columns) == 2
    assert table.cell(1, 1).text == "", "a missing cell is empty, not a dropped column"


def test_rows_still_render_when_cells_is_empty():
    table = _reopen(render_docx(_model(_block("table", rows=[["h"], ["v"]])))).tables[0]

    assert [table.cell(0, 0).text, table.cell(1, 0).text] == ["h", "v"]


# ---------------------------------------------------------------- the model's own sweep


def _sample(kind: str) -> tuple[Block, str]:
    """One block per kind plus the marker that proves it reached the document."""
    samples = {
        "heading": (_block("heading", text="H", level=3), "H"),
        "paragraph": (_block("paragraph", text="P"), "P"),
        "bullets": (_block("bullets", items=["B"]), "B"),
        "numbered": (_block("numbered", items=["N"]), "N"),
        "table": (_block("table", rows=[["T"]]), "T"),
        "image": (_block("image", artifact_slug="slug-1"), "[image: slug-1]"),
        "pagebreak": (_block("pagebreak"), 'w:type="page"'),
        "code": (_block("code", text="C"), "C"),
    }
    return samples[kind]


def _visible_text(doc) -> str:
    parts = [para.text for para in doc.paragraphs]
    parts += [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    return "\n".join(parts)


def test_every_declared_block_kind_still_renders_something():
    """The model says a writer MUST handle every kind; a skipped kind drops content."""
    assert len(BLOCK_KINDS) == 8, "vacuity floor: this sweep must not shrink silently"

    for kind in BLOCK_KINDS:
        block, marker = _sample(kind)
        data = render_docx(_model(block))
        haystack = _visible_text(_reopen(data)) + _document_xml(data).decode("utf-8")
        assert marker in haystack, f"{kind} rendered nothing"


def test_a_run_carrying_block_still_renders_something_for_every_kind_that_has_runs():
    for kind in ("heading", "paragraph", "code"):
        block = _block(kind, runs=[_Run("marker", bold=True)])
        assert "marker" in _visible_text(_reopen(render_docx(_model(block)))), kind


# ------------------------------------------------------------ through the REAL  reader


def test_a_run_carrying_document_re_reads_through_the_real_reader(tmp_path):
    """The repo's in-process validity proof: the same reader that ingests uploads.

    A `w:hyperlink` keeps its visible text out of `w:p`'s own `w:r` children, so a reader
    that walked only those runs would silently drop the link's words.
    """
    runs = [_Run("plain "), _Run("strong", bold=True), _Run(", see "), _Run("the docs", link=_URL)]
    path = tmp_path / "out.docx"
    path.write_bytes(render_docx(_model(_block("paragraph", runs=runs), title="Doc")))

    text, meta = FileReader().read(str(path))

    assert meta["format"] == "docx"
    assert "plain strong, see the docs" in text, "the link's words must survive the reader"


# ------------------------------------------------------ soul guardrail 5: writers  pure

#: A writer takes a model and returns bytes. Anything below would make it reach outward.
_IMPURE = (
    "open(",
    "requests",
    "httpx",
    "urllib",
    "socket",
    "subprocess",
    "config_dir",
    "get_config",
    "load_config",
    "os.environ",
    "getenv",
    "Path(",
)


def _impurities(source: str) -> list[str]:
    return sorted(token for token in _IMPURE if token in source)


def test_the_docx_writer_reaches_for_nothing():
    source = Path(docx_writer.__file__).read_text(encoding="utf-8")

    assert _impurities(source) == [], "a writer takes a model and returns bytes"


def test_the_purity_matcher_fires_when_the_impurity_is_present():
    """Vacuity floor: a rail that matches nothing asserts nothing."""
    impure = "import requests\nimport os\nwith open('x') as fh:\n    os.environ['A']\n"

    assert _impurities(impure) == ["open(", "os.environ", "requests"]
