"""DFE-4 — the binary artifact write path and the document model read/render pair.

Three routes land here: ``PUT /api/artifacts/{slug}/raw`` (bytes in),
``GET /api/artifacts/{slug}/model`` (structure out) and
``PUT /api/artifacts/{slug}/model`` (structure in, re-rendered server-side). This suite
asserts each of the atom's six clauses **at the route**, not at the mechanism: a test
proving :func:`gideon.http_errors.json_error` works, or that
``update_binary(expect_version=…)`` raises, would say nothing about whether these
handlers call them.

**Two clauses need more than a status code, and both are written that way.**

*"An oversized body is refused BEFORE buffering."* A test that posts a big body and
reads a 413 cannot tell a pre-check from a post-check — both answer 413, and the
post-check has already spent the memory. So the load-bearing test drives the handler
with a request whose body **raises if anything reads it**
(:class:`_ExplodingPayload`): the refusal can only be produced from ``Content-Length``,
because reading is fatal. A second test then proves the same refusal fires through the
real router with a real body, so the mocked-request test is not the only evidence.

*"Exactly ONE version, ONE SEL row."* Both are counted, never merely present: two
versions per save silently doubles a document's history, and two audit rows split one
action's trail across two entries. ``== 1``, not ``>= 1``.

Every refusal test has a matching acceptance so no guard can pass by refusing
everything, and the SEL these tests read is ``checks/runtime/conftest.py``'s per-test temp log —
a leak into the real ``~/.gideon`` fails conftest's own rail.
"""

from __future__ import annotations

import io
import json
from unittest.mock import MagicMock, patch

import pytest
from aiohttp import web
from aiohttp.test_utils import TestClient, TestServer, make_mocked_request

from gideon.security.sel import sel
from gideon.workspace.artifacts import handlers as handlers_mod
from gideon.workspace.artifacts import registry
from gideon.workspace.artifacts.handlers import register_artifact_routes
from gideon.workspace.artifacts.models import (
    MAX_BINARY_CONTENT_BYTES,
    ArtifactVersionConflict,
    mime_for_ext,
)
from gideon.workspace.artifacts.native import NativeArtifactProvider
from gideon.workspace.documents.deck_json import deck_to_dict
from gideon.workspace.documents.docx_parser import parse_docx
from gideon.workspace.documents.model import (
    Block,
    Bullet,
    DeckModel,
    DocumentModel,
    Run,
    Sheet,
    SheetCell,
    SheetModel,
    Slide,
)
from gideon.workspace.documents.model_json import document_from_dict, document_to_dict
from gideon.workspace.documents.sheet_json import sheet_to_dict
from gideon.workspace.documents.writers.docx_writer import render_docx
from gideon.workspace.documents.writers.pptx_writer import render_pptx
from gideon.workspace.documents.writers.xlsx_writer import render_xlsx

DOCX_MIME = mime_for_ext("docx")
PDF_MIME = mime_for_ext("pdf")
PNG_MIME = mime_for_ext("png")
XLSX_MIME = mime_for_ext("xlsx")
PPTX_MIME = mime_for_ext("pptx")

_OOXML_TELLS = (b"PK\x03\x04", b"word/document.xml", b"<w:")


def _authored(bold_word: str = "bold") -> DocumentModel:
    """A small document with one bold run — the thing an edit has to preserve."""
    return DocumentModel(
        title="Fidelity",
        blocks=[
            Block(kind="heading", text="Overview", level=1),
            Block(
                kind="paragraph",
                runs=[
                    Run(text="a "),
                    Run(text=bold_word, bold=True),
                    Run(text=" word"),
                ],
            ),
        ],
    )


def _docx_bytes(model: DocumentModel | None = None) -> bytes:
    """One render of *model*, and it IS lossless when parsed back.

    It was not before DFE-6: python-docx's default template ships asymmetric page margins
    (1.00in top/bottom, 1.25in left/right) which a single ``margin_in`` could not hold, so
    a document this repo generated itself reported a ``page_property`` loss. ``PageSetup``
    now carries a margin per edge, so that geometry is representable and the report on a
    freshly generated document is empty."""
    return render_docx(model if model is not None else _authored())


def _settled_docx_bytes(model: DocumentModel | None = None) -> bytes:
    """*model* rendered, parsed and re-rendered, so its page setup is EXPLICIT in the model
    rather than inherited from the template.

    Since DFE-6 one render already parses losslessly, so this no longer exists to dodge a
    margin loss — it exists for the tests that need the page setup to be a stated fact of
    the model they hold."""
    parsed, _ = parse_docx(_docx_bytes(model))
    return render_docx(parsed)


def _lossy_docx_bytes() -> bytes:
    """A document carrying a construct the model genuinely CANNOT hold: a two-paragraph
    header, against ``PageSetup.header_text``'s one string.

    Until DFE-6 the vehicle here was the template's asymmetric page margins. Those are
    representable now, so a test that still used them would assert `lossless is False`
    about a lossless document — the vacuity leg needs a construct that is honestly
    unrepresentable, not one that used to be.
    """
    from docx import Document as _Docx

    doc = _Docx(io.BytesIO(_settled_docx_bytes()))
    header = doc.sections[0].header
    header.paragraphs[0].text = "line one"
    header.add_paragraph("line two")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture
def provider(tmp_path) -> NativeArtifactProvider:
    return NativeArtifactProvider(root=tmp_path / "artifacts")


@pytest.fixture
def document_editing_on(monkeypatch):
    """``dashboard.document_editing`` ON — the §C6 consent gate ``PUT …/model`` reads.

    Patched rather than assumed, for two reasons. The flag is OFF by default, so without
    it every model write here would be measuring the CONSENT gate instead of the contract
    under test; and ``AppConfig.load()`` inside the handler would otherwise read the
    developer's REAL ``config.json``, making this suite's result depend on their settings.
    A default ``AppConfig`` with one field flipped keeps the patch honest — it is the same
    object the route reads in production, not a mock that says yes to everything.
    (The refusal itself is covered in ``checks/runtime/test_document_editing_gate.py``.)
    """
    from gideon.core.config import AppConfig

    def _load() -> AppConfig:
        cfg = AppConfig()
        cfg.dashboard.document_editing = True
        return cfg

    monkeypatch.setattr(AppConfig, "load", staticmethod(_load))


@pytest.fixture
def patched_native(provider, document_editing_on):
    with patch.object(registry, "get_provider", return_value=provider):
        yield provider


def _state() -> MagicMock:
    state = MagicMock()
    state._restricted_keys = set()
    state._sessions = {}
    return state


def _app(state: MagicMock | None = None) -> web.Application:
    app = web.Application(client_max_size=MAX_BINARY_CONTENT_BYTES * 2)
    app["state"] = state if state is not None else _state()
    register_artifact_routes(app)
    return app


async def _client(state: MagicMock | None = None) -> TestClient:
    client = TestClient(TestServer(_app(state)))
    await client.start_server()
    return client


def _docx_artifact(provider, *, name: str = "Report", data: bytes | None = None):
    return provider.create_binary(
        name=name,
        data=data if data is not None else _settled_docx_bytes(),
        mime=DOCX_MIME,
        kind="docx",
        actor="agent",
    )


def _sel_rows(operation: str) -> list[dict]:
    return [e for e in sel().recent(limit=500) if e.get("operation") == operation]


@pytest.mark.asyncio
async def test_a_stale_if_match_is_refused_409_and_nothing_is_written(
    patched_native,
) -> None:
    prov = patched_native
    art = _docx_artifact(prov)
    prov.update_binary(art.slug, data=_docx_bytes(_authored("second")), actor="agent")
    before = prov.get(art.slug)
    assert before.version == 2
    original = prov.raw_bytes(art.slug)[0]

    client = await _client()
    try:
        resp = await client.put(
            f"/api/artifacts/{art.slug}/raw",
            data=_docx_bytes(_authored("third")),
            headers={"Content-Type": DOCX_MIME, "If-Match": "1"},
        )
        assert resp.status == 409
        body = await resp.json()
        assert body["error"]["code"] == "version_conflict"
        assert body["error"]["expected"] == 2 and body["error"]["supplied"] == 1
    finally:
        await client.close()

    after = prov.get(art.slug)
    assert after.version == 2, "a refused write must not bump a version"
    assert (
        prov.raw_bytes(art.slug)[0] == original
    ), "a refused write must not touch the body"


@pytest.mark.asyncio
async def test_a_missing_if_match_is_refused_428(patched_native) -> None:
    """The precondition is REQUIRED, and 'you did not send it' is its own remedy: add the
    header. Conflating it with 409 would tell a client to reload when its version is fine.
    """
    art = _docx_artifact(patched_native)
    client = await _client()
    try:
        resp = await client.put(
            f"/api/artifacts/{art.slug}/raw",
            data=_docx_bytes(),
            headers={"Content-Type": DOCX_MIME},
        )
        assert resp.status == 428
        body = await resp.json()
        assert body["error"]["code"] == "if_match_required"
        assert body["error"]["version"] == 1
    finally:
        await client.close()
    assert patched_native.get(art.slug).version == 1


@pytest.mark.asyncio
async def test_a_malformed_if_match_is_refused_400(patched_native) -> None:
    art = _docx_artifact(patched_native)
    client = await _client()
    try:
        resp = await client.put(
            f"/api/artifacts/{art.slug}/raw",
            data=_docx_bytes(),
            headers={"Content-Type": DOCX_MIME, "If-Match": 'W/"abc"'},
        )
        assert resp.status == 400
        assert (await resp.json())["error"]["code"] == "if_match_malformed"
    finally:
        await client.close()


@pytest.mark.parametrize("header", ["1", '"1"', 'W/"1"'])
@pytest.mark.asyncio
async def test_a_current_if_match_is_accepted(patched_native, header) -> None:
    """VACUITY for clause 1: the guard can pass. Quoted and weak forms are tolerated
    because HTTP clients add them unasked — a 409 for a correct version would be worse
    than no precondition at all."""
    prov = patched_native
    art = _docx_artifact(prov)
    client = await _client()
    try:
        resp = await client.put(
            f"/api/artifacts/{art.slug}/raw",
            data=_docx_bytes(_authored("edited")),
            headers={"Content-Type": DOCX_MIME, "If-Match": header},
        )
        assert resp.status == 200, await resp.text()
        assert (await resp.json())["version"] == 2
    finally:
        await client.close()


def test_the_precondition_is_enforced_inside_the_provider_lock(provider) -> None:
    """The handler's If-Match check is the well-worded refusal; THIS is the guarantee.

    A handler that only compared versions itself would leave a window between the read
    and the write — exactly the race a whole-document save exists to detect. So the
    comparison also happens under the provider's lock, and this asserts the provider
    refuses on its own.
    """
    art = _docx_artifact(provider)
    provider.update_binary(
        art.slug, data=_docx_bytes(_authored("second")), actor="agent"
    )
    with pytest.raises(ArtifactVersionConflict) as caught:
        provider.update_binary(art.slug, data=b"x", actor="user", expect_version=1)
    assert caught.value.expected == 2 and caught.value.supplied == 1
    assert (
        provider.update_binary(art.slug, data=_docx_bytes(), expect_version=2).version
        == 3
    )


class _ExplodingPayload:
    """A request body that raises the moment anything tries to read it.

    This is the instrument for clause 2. ``request.read()`` drains the payload through
    ``readany``; if the handler reaches it, the test fails with a message naming the
    defect instead of quietly passing on a 413 produced too late.
    """

    _BOOM = "the body was buffered before the size cap was checked"

    def __init__(self) -> None:
        self.reads = 0

    def set_read_chunk_size(self, size: int) -> None:
        self.reads += 1
        raise AssertionError(self._BOOM)

    async def readany(self) -> bytes:
        self.reads += 1
        raise AssertionError(self._BOOM)

    async def read(self, n: int = -1) -> bytes:
        self.reads += 1
        raise AssertionError(self._BOOM)


@pytest.mark.asyncio
async def test_an_oversized_body_is_refused_without_the_body_being_read(
    patched_native,
) -> None:
    art = _docx_artifact(patched_native)
    payload = _ExplodingPayload()
    over = MAX_BINARY_CONTENT_BYTES + 1
    request = make_mocked_request(
        "PUT",
        f"/api/artifacts/{art.slug}/raw",
        headers={
            "Content-Type": DOCX_MIME,
            "Content-Length": str(over),
            "If-Match": "1",
        },
        match_info={"slug": art.slug},
        payload=payload,
        app=_app(),
    )
    resp = await handlers_mod.api_artifact_raw_write(request)
    assert resp.status == 413
    body = json.loads(resp.body)
    assert body["error"]["code"] == "request_too_large"
    assert body["error"]["cap_bytes"] == MAX_BINARY_CONTENT_BYTES
    assert body["error"]["declared_bytes"] == over
    assert payload.reads == 0, "the handler touched the body before refusing"


@pytest.mark.asyncio
async def test_an_underlimit_body_does_reach_the_payload(patched_native) -> None:
    """VACUITY for the instrument above: an in-limit request DOES read the body, so
    ``reads == 0`` in the oversized case is evidence about the cap and not about a
    handler that never reads anything."""
    art = _docx_artifact(patched_native)
    payload = _ExplodingPayload()
    request = make_mocked_request(
        "PUT",
        f"/api/artifacts/{art.slug}/raw",
        headers={"Content-Type": DOCX_MIME, "Content-Length": "10", "If-Match": "1"},
        match_info={"slug": art.slug},
        payload=payload,
        app=_app(),
    )
    with pytest.raises(AssertionError, match="buffered"):
        await handlers_mod.api_artifact_raw_write(request)
    assert payload.reads == 1


@pytest.mark.asyncio
async def test_the_cap_also_refuses_through_the_real_router(
    patched_native, monkeypatch
) -> None:
    """The same refusal, driven end to end with a real body, at a shrunken cap — so the
    mocked-request test above is not the only evidence that the route enforces this."""
    monkeypatch.setattr(handlers_mod, "MAX_BINARY_CONTENT_BYTES", 64)
    prov = patched_native
    art = _docx_artifact(prov)
    client = await _client()
    try:
        resp = await client.put(
            f"/api/artifacts/{art.slug}/raw",
            data=b"x" * 512,
            headers={"Content-Type": DOCX_MIME, "If-Match": "1"},
        )
        assert resp.status == 413
        assert (await resp.json())["error"]["cap_bytes"] == 64
        ok = await client.put(
            f"/api/artifacts/{art.slug}/raw",
            data=b"y" * 32,
            headers={"Content-Type": DOCX_MIME, "If-Match": "1"},
        )
        assert ok.status == 200, await ok.text()
    finally:
        await client.close()
    assert prov.get(art.slug).version == 2


@pytest.mark.asyncio
async def test_a_chunked_body_is_refused_411(patched_native) -> None:
    """No ``Content-Length`` means nothing to check before reading, which is the one
    thing this route cannot do — so it refuses rather than falling back to reading."""
    art = _docx_artifact(patched_native)
    payload = _ExplodingPayload()
    request = make_mocked_request(
        "PUT",
        f"/api/artifacts/{art.slug}/raw",
        headers={
            "Content-Type": DOCX_MIME,
            "Transfer-Encoding": "chunked",
            "If-Match": "1",
        },
        match_info={"slug": art.slug},
        payload=payload,
        app=_app(),
    )
    resp = await handlers_mod.api_artifact_raw_write(request)
    assert resp.status == 411
    assert json.loads(resp.body)["error"]["code"] == "content_length_required"
    assert payload.reads == 0


@pytest.mark.asyncio
async def test_a_pdf_body_is_refused_for_a_docx_artifact(patched_native) -> None:
    prov = patched_native
    art = _docx_artifact(prov)
    original = prov.raw_bytes(art.slug)[0]
    client = await _client()
    try:
        resp = await client.put(
            f"/api/artifacts/{art.slug}/raw",
            data=b"%PDF-1.7 not really",
            headers={"Content-Type": PDF_MIME, "If-Match": "1"},
        )
        assert resp.status == 409
        body = await resp.json()
        assert body["error"]["code"] == "mime_kind_mismatch"
        assert (
            body["error"]["kind"] == "docx" and body["error"]["declared_kind"] == "pdf"
        )
        ok = await client.put(
            f"/api/artifacts/{art.slug}/raw",
            data=_docx_bytes(_authored("kept")),
            headers={"Content-Type": DOCX_MIME, "If-Match": "1"},
        )
        assert ok.status == 200, await ok.text()
    finally:
        await client.close()
    assert prov.raw_bytes(art.slug)[0] != original


@pytest.mark.asyncio
async def test_an_unstorable_content_type_is_refused_415(patched_native) -> None:
    art = _docx_artifact(patched_native)
    client = await _client()
    try:
        resp = await client.put(
            f"/api/artifacts/{art.slug}/raw",
            data=b"hello",
            headers={"Content-Type": "text/plain", "If-Match": "1"},
        )
        assert resp.status == 415
        assert (await resp.json())["error"]["code"] == "unsupported_media_type"
    finally:
        await client.close()


@pytest.mark.asyncio
async def test_a_text_kind_artifact_has_no_binary_body_to_replace(
    patched_native,
) -> None:
    prov = patched_native
    art = prov.create(name="Notes", content="# hi", kind="markdown")
    client = await _client()
    try:
        resp = await client.put(
            f"/api/artifacts/{art.slug}/raw",
            data=_docx_bytes(),
            headers={"Content-Type": DOCX_MIME, "If-Match": "1"},
        )
        assert resp.status == 409
        assert (await resp.json())["error"]["code"] == "kind_not_binary"
    finally:
        await client.close()


@pytest.mark.asyncio
async def test_an_image_artifact_takes_an_image_body(patched_native) -> None:
    """VACUITY for the kind guard across the whole binary family, not just docx."""
    prov = patched_native
    art = prov.create_binary(name="Cat", data=b"\x89PNG\r\n\x1a\nold", mime=PNG_MIME)
    client = await _client()
    try:
        resp = await client.put(
            f"/api/artifacts/{art.slug}/raw",
            data=b"\x89PNG\r\n\x1a\nnew",
            headers={"Content-Type": PNG_MIME, "If-Match": "1"},
        )
        assert resp.status == 200, await resp.text()
    finally:
        await client.close()
    assert prov.raw_bytes(art.slug)[0] == b"\x89PNG\r\n\x1a\nnew"


@pytest.mark.asyncio
async def test_a_missing_slug_is_404_and_a_restricted_session_is_403(
    patched_native,
) -> None:
    client = await _client()
    try:
        gone = await client.put(
            "/api/artifacts/no-such-thing/raw",
            data=_docx_bytes(),
            headers={"Content-Type": DOCX_MIME, "If-Match": "1"},
        )
        assert gone.status == 404
        assert (await gone.json())["error"]["code"] == "not_found"
    finally:
        await client.close()

    art = _docx_artifact(patched_native)
    with patch.object(handlers_mod, "_is_restricted_session", return_value=True):
        client = await _client()
        try:
            resp = await client.put(
                f"/api/artifacts/{art.slug}/raw",
                data=_docx_bytes(),
                headers={"Content-Type": DOCX_MIME, "If-Match": "1"},
            )
            assert resp.status == 403
            assert (await resp.json())["error"]["code"] == "forbidden"
        finally:
            await client.close()
    assert patched_native.get(art.slug).version == 1


@pytest.mark.asyncio
async def test_an_accepted_write_bumps_one_version_and_logs_one_audit_row(
    patched_native,
) -> None:
    prov = patched_native
    art = _docx_artifact(prov)
    versions_before = prov.list_versions(art.slug)
    rows_before = len(_sel_rows("artifact.raw_write"))
    new_bytes = _docx_bytes(_authored("edited"))

    client = await _client()
    try:
        resp = await client.put(
            f"/api/artifacts/{art.slug}/raw",
            data=new_bytes,
            headers={"Content-Type": DOCX_MIME, "If-Match": "1"},
        )
        assert resp.status == 200, await resp.text()
        assert await resp.json() == {"slug": art.slug, "version": 2, "mime": DOCX_MIME}
    finally:
        await client.close()

    versions_after = prov.list_versions(art.slug)
    assert (
        len(versions_after) - len(versions_before) == 1
    ), f"expected exactly one new version, got {versions_before} -> {versions_after}"
    assert prov.get(art.slug).version == 2
    assert prov.raw_bytes(art.slug)[0] == new_bytes
    assert prov.raw_bytes(art.slug, version=1)[0] != new_bytes

    rows = _sel_rows("artifact.raw_write")
    assert len(rows) - rows_before == 1, f"expected exactly one audit row, got {rows}"
    written = rows[-1]
    assert written["outcome"] == "ok"
    assert f"bytes={len(new_bytes)}" in written["resources"], written["resources"]
    assert f"slug={art.slug}" in written["resources"]
    assert "version=2" in written["resources"]


@pytest.mark.asyncio
async def test_a_refused_write_logs_a_denial_and_no_ok_row(patched_native) -> None:
    """VACUITY for the audit count: the operation is logged on the refusal path too, so
    ``== 1`` above is a count of ACCEPTED writes rather than of everything."""
    art = _docx_artifact(patched_native)
    before = len(_sel_rows("artifact.raw_write"))
    client = await _client()
    try:
        resp = await client.put(
            f"/api/artifacts/{art.slug}/raw",
            data=_docx_bytes(),
            headers={"Content-Type": PDF_MIME, "If-Match": "1"},
        )
        assert resp.status == 409
    finally:
        await client.close()
    rows = _sel_rows("artifact.raw_write")[before:]
    assert len(rows) == 1 and rows[0]["outcome"] == "denied"


@pytest.mark.asyncio
async def test_get_model_returns_the_parsed_model_and_the_loss_report(
    patched_native,
) -> None:
    art = _docx_artifact(patched_native)
    client = await _client()
    try:
        resp = await client.get(f"/api/artifacts/{art.slug}/model")
        assert resp.status == 200, await resp.text()
        body = await resp.json()
    finally:
        await client.close()

    assert body["slug"] == art.slug and body["kind"] == "docx" and body["version"] == 1
    assert body["model"]["title"] == "Fidelity"
    kinds = [b["kind"] for b in body["model"]["blocks"]]
    assert kinds == ["heading", "paragraph"]
    runs = body["model"]["blocks"][1]["runs"]
    assert [r["text"] for r in runs] == ["a ", "bold", " word"]
    assert [r["bold"] for r in runs] == [False, True, False]
    assert set(body["loss"]) == {"lossless", "kinds", "summary", "items"}
    assert body["loss"]["lossless"] is True
    assert body["loss"]["items"] == [] and body["loss"]["summary"] == "no losses"
    assert document_to_dict(document_from_dict(body["model"])) == body["model"]


@pytest.mark.asyncio
async def test_the_loss_report_names_what_did_not_fit(patched_native) -> None:
    """VACUITY for the clause above: ``lossless`` can be False, and when it is, the
    report names the construct and where it was. A report that only ever says "no
    losses" would let §C5's warning surface stay silent for a lossy document."""
    art = _docx_artifact(patched_native, data=_lossy_docx_bytes())
    client = await _client()
    try:
        body = await (await client.get(f"/api/artifacts/{art.slug}/model")).json()
    finally:
        await client.close()
    loss = body["loss"]
    assert loss["lossless"] is False
    assert loss["kinds"] == ["header_footer"]
    assert loss["summary"] == "header_footer×1"
    (item,) = loss["items"]
    assert item["kind"] == "header_footer" and item["where"] == "document"
    assert "2 non-empty paragraphs" in item["detail"]


@pytest.mark.asyncio
async def test_the_model_response_carries_no_ooxml(patched_native) -> None:
    """Clause 5's real claim: the editor's read is JSON structure, never document bytes."""
    art = _docx_artifact(patched_native)
    client = await _client()
    try:
        resp = await client.get(f"/api/artifacts/{art.slug}/model")
        assert resp.content_type == "application/json"
        raw = await resp.read()
    finally:
        await client.close()
    for tell in _OOXML_TELLS:
        assert tell not in raw, f"the model response leaked {tell!r}"
    stored = patched_native.raw_bytes(art.slug)[0]
    assert _OOXML_TELLS[0] in stored


@pytest.mark.asyncio
async def test_a_model_write_renders_server_side_and_the_edit_survives(
    patched_native,
) -> None:
    """The save half end to end: the client posts STRUCTURE, the server renders, and the
    formatting is still there when the stored bytes are parsed back."""
    prov = patched_native
    art = _docx_artifact(prov)
    client = await _client()
    try:
        loaded = await (await client.get(f"/api/artifacts/{art.slug}/model")).json()
        model = loaded["model"]
        model["blocks"][1]["runs"][2]["bold"] = True
        payload = json.dumps({"model": model}).encode()
        for tell in _OOXML_TELLS:
            assert tell not in payload, "the client had to construct OOXML to save"
        versions_before = prov.list_versions(art.slug)
        rows_before = len(_sel_rows("artifact.model_write"))
        resp = await client.put(
            f"/api/artifacts/{art.slug}/model",
            data=payload,
            headers={
                "Content-Type": "application/json",
                "If-Match": str(loaded["version"]),
            },
        )
        assert resp.status == 200, await resp.text()
        assert await resp.json() == {"slug": art.slug, "version": 2, "mime": DOCX_MIME}
    finally:
        await client.close()

    assert len(prov.list_versions(art.slug)) - len(versions_before) == 1
    rows = _sel_rows("artifact.model_write")
    assert len(rows) - rows_before == 1
    stored = prov.raw_bytes(art.slug)[0]
    assert stored.startswith(b"PK\x03\x04"), "the server stored real .docx bytes"
    reparsed, loss = parse_docx(stored)
    assert [(r.text, r.bold) for r in reparsed.blocks[1].runs] == [
        ("a ", False),
        ("bold word", True),
    ]
    assert loss.lossless, loss.summary()


@pytest.mark.asyncio
async def test_a_model_write_honors_the_same_stale_if_match_refusal(
    patched_native,
) -> None:
    prov = patched_native
    art = _docx_artifact(prov)
    prov.update_binary(art.slug, data=_docx_bytes(_authored("second")), actor="agent")
    original = prov.raw_bytes(art.slug)[0]
    client = await _client()
    try:
        resp = await client.put(
            f"/api/artifacts/{art.slug}/model",
            json={"model": document_to_dict(_authored("third"))},
            headers={"If-Match": "1"},
        )
        assert resp.status == 409
        assert (await resp.json())["error"]["code"] == "version_conflict"
        missing = await client.put(
            f"/api/artifacts/{art.slug}/model",
            json={"model": document_to_dict(_authored("third"))},
        )
        assert missing.status == 428
    finally:
        await client.close()
    assert prov.get(art.slug).version == 2
    assert prov.raw_bytes(art.slug)[0] == original


@pytest.mark.asyncio
async def test_a_bad_model_is_refused_and_nothing_is_written(patched_native) -> None:
    prov = patched_native
    art = _docx_artifact(prov)
    original = prov.raw_bytes(art.slug)[0]
    client = await _client()
    try:
        for bad, code in (
            ({"model": {"title": "x", "typo": 1}}, "invalid_model"),
            ({"model": {"blocks": [{"kind": "nonsense"}]}}, "invalid_model"),
            (
                {"model": {"blocks": [{"kind": "paragraph", "text": 7}]}},
                "invalid_model",
            ),
            ({"model": "a string"}, "invalid_model"),
            ({}, "invalid_model"),
            ([1, 2, 3], "invalid_body"),
        ):
            resp = await client.put(
                f"/api/artifacts/{art.slug}/model",
                json=bad,
                headers={"If-Match": "1"},
            )
            assert resp.status == 400, (bad, await resp.text())
            assert (await resp.json())["error"]["code"] == code, bad
        broken = await client.put(
            f"/api/artifacts/{art.slug}/model",
            data=b"{not json",
            headers={"Content-Type": "application/json", "If-Match": "1"},
        )
        assert broken.status == 400
        assert (await broken.json())["error"]["code"] == "invalid_json"
    finally:
        await client.close()
    assert prov.get(art.slug).version == 1
    assert prov.raw_bytes(art.slug)[0] == original


@pytest.mark.asyncio
async def test_a_kind_without_a_shipped_parser_has_no_model(patched_native) -> None:
    """An image has a binary body but no document model, and saying so is the honest
    answer — an empty model would read to the editor as an empty document."""
    prov = patched_native
    art = prov.create_binary(name="Cat", data=b"\x89PNG\r\n\x1a\n", mime=PNG_MIME)
    client = await _client()
    try:
        resp = await client.get(f"/api/artifacts/{art.slug}/model")
        assert resp.status == 415
        body = await resp.json()
        assert body["error"]["code"] == "model_unavailable"
        assert body["error"]["model_kinds"] == ["docx", "xlsx", "pptx"]
        write = await client.put(
            f"/api/artifacts/{art.slug}/model",
            json={"model": document_to_dict(_authored())},
            headers={"If-Match": "1"},
        )
        assert write.status == 415
    finally:
        await client.close()


@pytest.mark.asyncio
async def test_unparseable_stored_bytes_are_a_400_not_a_500(patched_native) -> None:
    prov = patched_native
    art = _docx_artifact(prov, data=b"not a zip at all")
    client = await _client()
    try:
        resp = await client.get(f"/api/artifacts/{art.slug}/model")
        assert resp.status == 400
        assert (await resp.json())["error"]["code"] == "model_parse_failed"
    finally:
        await client.close()


@pytest.mark.asyncio
async def test_a_restricted_session_cannot_save_a_model(patched_native) -> None:
    art = _docx_artifact(patched_native)
    with patch.object(handlers_mod, "_is_restricted_session", return_value=True):
        client = await _client()
        try:
            resp = await client.put(
                f"/api/artifacts/{art.slug}/model",
                json={"model": document_to_dict(_authored())},
                headers={"If-Match": "1"},
            )
            assert resp.status == 403
        finally:
            await client.close()
    assert patched_native.get(art.slug).version == 1


def test_the_new_routes_are_registered_and_ordered_after_the_literal_paths() -> None:
    """The atom's routes exist on the real table, and the literal-path-before-``{slug}``
    rule the surrounding comments explain still holds — a ``{slug}`` pattern registered
    ahead of ``pinned``/``folders``/``deployed`` would swallow them."""
    app = _app()
    registered = [
        (r.method, r.get_info().get("path") or r.get_info().get("formatter"))
        for r in app.router.routes()
    ]
    for wanted in (
        ("PUT", "/api/artifacts/{slug}/raw"),
        ("GET", "/api/artifacts/{slug}/model"),
        ("PUT", "/api/artifacts/{slug}/model"),
    ):
        assert wanted in registered, f"{wanted} is not on the route table"
    paths = [path for _, path in registered]
    first_slug = min(i for i, p in enumerate(paths) if p and "{slug}" in p)
    for literal in (
        "/api/artifacts/pinned",
        "/api/artifacts/deployed",
        "/api/artifacts/folders",
    ):
        assert paths.index(literal) < first_slug


def test_only_the_download_route_serves_document_bytes() -> None:
    """Clause 5's structural half. ``GET …/raw`` is the DOWNLOAD affordance and serves
    the real bytes on purpose; the editor's three-route circuit does not. Asserted by
    reading which handlers can return a non-JSON body.
    """
    import inspect

    for name in (
        "api_artifact_raw_write",
        "api_artifact_model",
        "api_artifact_model_write",
    ):
        source = inspect.getsource(getattr(handlers_mod, name))
        assert "web.Response(" not in source, f"{name} can return a raw body"
    assert "web.Response(" in inspect.getsource(handlers_mod.api_artifact_raw)


def _sheet_bytes() -> bytes:
    """A workbook with a formula, a number format and a bold header."""
    return render_xlsx(
        SheetModel(
            sheets=[
                Sheet(
                    name="Q1",
                    cells=[
                        [
                            SheetCell(value="Item", bold=True),
                            SheetCell(value="Qty", bold=True),
                        ],
                        [SheetCell(value="Pens"), SheetCell(value=12)],
                        [SheetCell(value="Total"), SheetCell(formula="=SUM(B2:B2)")],
                    ],
                )
            ]
        )
    )


def _xlsx_artifact(provider):
    return provider.create_binary(
        name="Budget", data=_sheet_bytes(), mime=XLSX_MIME, kind="xlsx", actor="agent"
    )


@pytest.mark.asyncio
async def test_a_spreadsheet_serves_its_model_through_the_same_route(
    patched_native,
) -> None:
    """``GET …/model`` must resolve the parser BY KIND. Before DFE-7 this route answered
    415 for every .xlsx, so this is the call site the codec table exists for."""
    prov = patched_native
    art = _xlsx_artifact(prov)
    client = await _client()
    try:
        resp = await client.get(f"/api/artifacts/{art.slug}/model")
        assert resp.status == 200
        body = await resp.json()
        assert body["kind"] == "xlsx"
        cells = body["model"]["sheets"][0]["cells"]
        assert cells[2][1]["formula"] == "=SUM(B2:B2)"
        assert cells[2][1]["value"] is None
        assert "blocks" not in body["model"]
        assert body["loss"]["kinds"] == ["formula_cached_value"]
    finally:
        await client.close()


@pytest.mark.asyncio
async def test_a_spreadsheet_cell_edit_survives_the_save_and_read_back(
    patched_native,
) -> None:
    """The full editor circuit at the route: GET the model, change one cell, PUT it back,
    then read the STORED BYTES with openpyxl — not our parser — and find the edit."""
    from openpyxl import load_workbook

    prov = patched_native
    art = _xlsx_artifact(prov)
    client = await _client()
    try:
        loaded = await (await client.get(f"/api/artifacts/{art.slug}/model")).json()
        model = loaded["model"]
        model["sheets"][0]["cells"][1][0]["value"] = "Notebooks"
        model["sheets"][0]["cells"][1][1]["number_format"] = "#,##0.00"
        resp = await client.put(
            f"/api/artifacts/{art.slug}/model",
            json={"model": model},
            headers={"If-Match": str(loaded["version"])},
        )
        assert resp.status == 200, await resp.text()

        stored, _ = prov.raw_bytes(art.slug)
        ws = load_workbook(io.BytesIO(stored))["Q1"]
        assert ws["A2"].value == "Notebooks"
        assert ws["B2"].number_format == "#,##0.00"
        assert ws["B3"].data_type == "f"
        assert ws["B3"].value == "=SUM(B2:B2)"
        assert ws["A1"].font.bold is True
    finally:
        await client.close()


@pytest.mark.asyncio
async def test_a_spreadsheet_save_is_refused_while_document_editing_is_off(
    provider, monkeypatch
) -> None:
    """The §C6 consent gate governs the sheet editor too — it is one route, so a second
    enforcement point was never written, and this proves the first one covers xlsx."""
    from gideon.core.config import AppConfig

    monkeypatch.setattr(AppConfig, "load", staticmethod(AppConfig))
    with patch.object(registry, "get_provider", return_value=provider):
        art = _xlsx_artifact(provider)
        client = await _client()
        try:
            resp = await client.put(
                f"/api/artifacts/{art.slug}/model",
                json={"model": sheet_to_dict(SheetModel.from_rows({"S": [["x"]]}))},
                headers={"If-Match": "1"},
            )
            assert resp.status == 403
            assert (await resp.json())["error"]["code"] == "document_editing_off"
        finally:
            await client.close()


@pytest.mark.asyncio
async def test_a_stale_spreadsheet_save_is_refused_not_merged(patched_native) -> None:
    """Two tabs over one workbook collide with a 409, exactly as they do for a document —
    the guard chain is shared, and this is the assertion that it really is shared."""
    prov = patched_native
    art = _xlsx_artifact(prov)
    client = await _client()
    try:
        loaded = await (await client.get(f"/api/artifacts/{art.slug}/model")).json()
        first = await client.put(
            f"/api/artifacts/{art.slug}/model",
            json={"model": loaded["model"]},
            headers={"If-Match": str(loaded["version"])},
        )
        assert first.status == 200
        second = await client.put(
            f"/api/artifacts/{art.slug}/model",
            json={"model": loaded["model"]},
            headers={"If-Match": str(loaded["version"])},
        )
        assert second.status == 409
    finally:
        await client.close()


@pytest.mark.asyncio
async def test_a_malformed_spreadsheet_model_is_refused_before_any_write(
    patched_native,
) -> None:
    """Strict deserialization at the route: an unknown field is a 400 and the artifact is
    untouched, so a client that misspelled a format learns it rather than losing it."""
    prov = patched_native
    art = _xlsx_artifact(prov)
    original = prov.raw_bytes(art.slug)[0]
    client = await _client()
    try:
        payload = sheet_to_dict(SheetModel.from_rows({"S": [["x"]]}))
        payload["sheets"][0]["cells"][0][0]["numberformat"] = "0.0%"
        resp = await client.put(
            f"/api/artifacts/{art.slug}/model",
            json={"model": payload},
            headers={"If-Match": "1"},
        )
        assert resp.status == 400
        assert (await resp.json())["error"]["code"] == "invalid_model"
        assert prov.get(art.slug).version == 1
        assert prov.raw_bytes(art.slug)[0] == original
    finally:
        await client.close()


def _deck_bytes() -> bytes:
    """A deck with a nested outline, speaker notes and a cover slide."""
    return render_pptx(
        DeckModel(
            title="Quarterly Review",
            slides=[
                Slide(
                    title="Pipeline",
                    bullets=[
                        Bullet(text="Enterprise"),
                        Bullet(text="Two renewals at risk", level=1),
                        Bullet(text="Self-serve"),
                    ],
                    notes="lead with the renewals",
                )
            ],
        )
    )


def _pptx_artifact(provider):
    return provider.create_binary(
        name="Review", data=_deck_bytes(), mime=PPTX_MIME, kind="pptx", actor="agent"
    )


@pytest.mark.asyncio
async def test_a_deck_serves_its_model_through_the_same_route(patched_native) -> None:
    """``GET …/model`` resolves the parser BY KIND. Before DFE-8 this route answered 415
    for every .pptx, so the deck editor had nothing to load."""
    prov = patched_native
    art = _pptx_artifact(prov)
    client = await _client()
    try:
        resp = await client.get(f"/api/artifacts/{art.slug}/model")
        assert resp.status == 200
        body = await resp.json()
        assert body["kind"] == "pptx"
        model = body["model"]
        assert model["title"] == "Quarterly Review"
        assert [(b["text"], b["level"]) for b in model["slides"][0]["bullets"]] == [
            ("Enterprise", 0),
            ("Two renewals at risk", 1),
            ("Self-serve", 0),
        ]
        assert "blocks" not in model and "sheets" not in model
        assert body["loss"]["lossless"] is True
    finally:
        await client.close()


@pytest.mark.asyncio
async def test_a_slide_edit_survives_the_save_and_read_back(patched_native) -> None:
    """The atom's second clause, at the route: GET the model, edit a slide, PUT it back,
    then read the STORED BYTES with python-pptx — not with our parser — and find the edit
    AND the bullet depth that was never touched."""
    from pptx import Presentation

    from gideon.workspace.documents.pptx_shapes import body_placeholder

    prov = patched_native
    art = _pptx_artifact(prov)
    client = await _client()
    try:
        loaded = await (await client.get(f"/api/artifacts/{art.slug}/model")).json()
        model = loaded["model"]
        model["slides"][0]["title"] = "Pipeline — revised"
        model["slides"][0]["bullets"][1]["text"] = "One renewal at risk"
        model["slides"][0]["bullets"].append({"text": "Nested detail", "level": 2})
        resp = await client.put(
            f"/api/artifacts/{art.slug}/model",
            json={"model": model},
            headers={"If-Match": str(loaded["version"])},
        )
        assert resp.status == 200, await resp.text()

        stored, _ = prov.raw_bytes(art.slug)
        slides = list(Presentation(io.BytesIO(stored)).slides)
        assert slides[0].shapes.title.text == "Quarterly Review"
        assert slides[1].shapes.title.text == "Pipeline — revised"
        frame = body_placeholder(slides[1]).text_frame
        assert [(p.text, p.level) for p in frame.paragraphs] == [
            ("Enterprise", 0),
            ("One renewal at risk", 1),
            ("Self-serve", 0),
            ("Nested detail", 2),
        ]
        assert "lead with the renewals" in slides[1].notes_slide.notes_text_frame.text
    finally:
        await client.close()


@pytest.mark.asyncio
async def test_a_deck_save_is_refused_while_document_editing_is_off(
    provider, monkeypatch
) -> None:
    """The §C6 consent gate governs the deck editor too — one route, one enforcement
    point, and this is the proof it covers pptx rather than a second copy of the check.
    """
    from gideon.core.config import AppConfig

    monkeypatch.setattr(AppConfig, "load", staticmethod(AppConfig))
    with patch.object(registry, "get_provider", return_value=provider):
        art = _pptx_artifact(provider)
        client = await _client()
        try:
            resp = await client.put(
                f"/api/artifacts/{art.slug}/model",
                json={
                    "model": deck_to_dict(DeckModel(slides=[Slide.outline("A", ["x"])]))
                },
                headers={"If-Match": "1"},
            )
            assert resp.status == 403
            assert (await resp.json())["error"]["code"] == "document_editing_off"
        finally:
            await client.close()


@pytest.mark.asyncio
async def test_a_stale_deck_save_is_refused_not_merged(patched_native) -> None:
    """Two tabs over one deck collide with a 409, exactly as they do for a document and a
    workbook — the guard chain is shared, and this asserts that it really is shared."""
    prov = patched_native
    art = _pptx_artifact(prov)
    client = await _client()
    try:
        loaded = await (await client.get(f"/api/artifacts/{art.slug}/model")).json()
        first = await client.put(
            f"/api/artifacts/{art.slug}/model",
            json={"model": loaded["model"]},
            headers={"If-Match": str(loaded["version"])},
        )
        assert first.status == 200
        second = await client.put(
            f"/api/artifacts/{art.slug}/model",
            json={"model": loaded["model"]},
            headers={"If-Match": str(loaded["version"])},
        )
        assert second.status == 409
    finally:
        await client.close()


@pytest.mark.asyncio
async def test_a_deck_naming_an_unknown_layout_is_refused_before_any_write(
    patched_native,
) -> None:
    """Strict deserialization at the route: a layout the shipped template does not have is
    a 400 and the artifact is untouched, rather than a deck that silently re-laid itself
    out and told the user by opening wrong."""
    prov = patched_native
    art = _pptx_artifact(prov)
    original = prov.raw_bytes(art.slug)[0]
    client = await _client()
    try:
        payload = deck_to_dict(DeckModel(slides=[Slide.outline("A", ["x"])]))
        payload["slides"][0]["layout"] = "Titel Slide"
        resp = await client.put(
            f"/api/artifacts/{art.slug}/model",
            json={"model": payload},
            headers={"If-Match": "1"},
        )
        assert resp.status == 400
        assert (await resp.json())["error"]["code"] == "invalid_model"
        assert prov.get(art.slug).version == 1
        assert prov.raw_bytes(art.slug)[0] == original
    finally:
        await client.close()
