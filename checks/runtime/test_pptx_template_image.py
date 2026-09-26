"""Real template masters and canonical image bytes survive a deck edit round trip."""

import io

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

from gideon.workspace.artifacts import registry
from gideon.workspace.artifacts.native import NativeArtifactProvider
from gideon.workspace.documents.model import Bullet, DeckModel, Slide
from gideon.workspace.documents.pptx_parser import parse_pptx
from gideon.workspace.documents.writers.pptx_writer import render_pptx


def test_template_and_image_are_real_pptx_content(tmp_path, monkeypatch):
    provider = NativeArtifactProvider(tmp_path / "artifacts")
    monkeypatch.setattr(registry, "get_provider", lambda: provider)

    image_data = io.BytesIO()
    Image.new("RGB", (320, 200), "#3775a9").save(image_data, format="PNG")
    image = provider.create_binary(name="Evidence figure", kind="image", data=image_data.getvalue(), mime="image/png")

    template = Presentation()
    template.slide_width = Inches(11)
    template.slides.add_slide(template.slide_layouts[0]).shapes.title.text = "Example content"
    template_data = io.BytesIO()
    template.save(template_data)
    source = provider.create_binary(name="Client template", kind="pptx", data=template_data.getvalue(),
                                    mime="application/vnd.openxmlformats-officedocument.presentationml.presentation")

    model = DeckModel(slides=[Slide(title="Finding", bullets=[Bullet("Evidence")],
                                    notes="Source: https://example.org/report", artifact_slug=image.slug)],
                      template_slug=source.slug)
    rendered = render_pptx(model)
    pptx = Presentation(io.BytesIO(rendered))
    assert len(pptx.slides) == 1
    assert pptx.slide_width == Inches(11)
    assert any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in pptx.slides[0].shapes)

    reopened, _ = parse_pptx(rendered)
    assert (reopened.template_slug, reopened.template_version) == (source.slug, source.version)
    assert reopened.slides[0].artifact_slug == image.slug
    assert "https://example.org/report" in reopened.slides[0].notes
    assert any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in Presentation(io.BytesIO(render_pptx(reopened))).slides[0].shapes)
