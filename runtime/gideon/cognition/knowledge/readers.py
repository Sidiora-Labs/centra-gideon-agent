"""Text and document decoding with structured PDF layout capture."""

import os
import re
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Protocol, runtime_checkable

from gideon.security.security import is_sensitive_path

try:
    import pdfplumber
except ImportError:
    pdfplumber = None  # type: ignore[assignment]

try:
    from pptx import Presentation  # type: ignore[import-untyped]
except ImportError:
    Presentation = None  # type: ignore[assignment,misc]

try:
    from docx import Document  # type: ignore[import-untyped]
except ImportError:
    Document = None  # type: ignore[assignment,misc]

try:
    import html2text as _html2text_mod
except ImportError:
    _html2text_mod = None  # type: ignore[assignment]

try:
    from openpyxl import load_workbook as _load_workbook  # type: ignore[import-untyped]
except ImportError:
    _load_workbook = None  # type: ignore[assignment]


def _table_lines(rows, *, escape=False, pad=False):
    if not rows:
        return []
    width = max(map(len, rows)) if pad else None

    def render(row):
        cells = (row + [""] * (width - len(row)))[:width] if width is not None else row
        if escape:
            cells = [value.replace("|", "\\|") for value in cells]
        return "| " + " | ".join(cells) + " |"

    header = render(rows[0])
    separator = render(["---"] * (width if width is not None else len(rows[0])))
    return [header, separator, *(render(row) for row in rows[1:])]


def _render_docx_table(table) -> list[str]:
    rows = []
    for row in table.rows:
        cells: list = []
        for cell in row.cells:
            value = " ".join(cell.text.split()).replace("|", "\\|")
            if not (value and cells and value == cells[-1]):
                cells.append(value)
        if any(cells):
            rows.append(cells)
    return _table_lines(rows, pad=True)


def _read_error(error, label="file"):
    return f"Error reading {label}: {error}", {"format": "error", "error": str(error)}


def _missing_reader(format_name, package):
    message = f"{format_name} support requires {package}"
    return f"{message}: pip install {package}", {"format": "error", "error": message}


def _text_file(path, encoding):
    with open(path, "r", encoding=encoding) as stream:
        return stream.read()


@runtime_checkable
class OcrProvider(Protocol):
    """Text extraction boundary shared by image and scanned-PDF ingestion."""

    def ocr(self, image_path: str, *, page_number: int = 1) -> str: ...


@dataclass(frozen=True)
class PdfResourceLimits:
    max_file_bytes: int = 100 * 1024 * 1024
    max_pages: int = 500
    max_raster_pages: int = 50
    max_page_pixels: int = 20_000_000
    max_raster_pixels: int = 100_000_000
    max_raster_bytes: int = 100 * 1024 * 1024
    max_text_chars: int = 10_000_000
    raster_dpi: int = 150


class PdfResourceLimitError(ValueError):
    def __init__(self, resource: str, actual: int, limit: int):
        self.resource = resource
        self.actual = actual
        self.limit = limit
        super().__init__(f"PDF {resource} limit exceeded ({actual} > {limit})")


class FileReader:
    SUPPORTED = {
        "",
        ".md",
        ".markdown",
        ".txt",
        ".text",
        ".py",
        ".java",
        ".ts",
        ".js",
        ".rs",
        ".go",
        ".html",
        ".htm",
        ".docx",
        ".csv",
        ".tsv",
        ".log",
        ".json",
        ".yaml",
        ".yml",
        ".sh",
        ".rb",
        ".c",
        ".cpp",
        ".h",
        ".xlsx",
        ".xls",
    }
    _DISPATCH = {
        ".pdf": "_read_pdf",
        ".pptx": "_read_pptx",
        ".docx": "_read_docx",
        ".html": "_read_html",
        ".htm": "_read_html",
        ".xlsx": "_read_xlsx",
        ".xls": "_read_xlsx",
        ".csv": "_read_csv",
        ".tsv": "_read_csv",
    }
    _CSV_MAX_TABLE_ROWS = 500

    def __init__(
        self,
        *,
        ocr_provider: OcrProvider | None = None,
        pdf_limits: PdfResourceLimits | None = None,
    ) -> None:
        self.ocr_provider = ocr_provider
        self.pdf_limits = pdf_limits or PdfResourceLimits()

    def read(self, path: str) -> tuple[str, dict]:
        if is_sensitive_path(path):
            raise PermissionError(f"Refusing to read sensitive path: {path}")
        file = Path(path)
        extension = file.suffix.lower()
        metadata = dict(
            format=extension.lstrip("."),
            title=file.stem,
            file_size=os.path.getsize(path),
            extension=extension,
        )
        handler = self._DISPATCH.get(extension)
        body, details = (
            getattr(self, handler)(path)
            if handler
            else self._read_text(path, extension.lstrip("."))
        )
        metadata.update(details)
        metadata["line_count"] = 1 + body.count("\n") if body else 0
        return body, metadata

    def _read_text(self, path: str, fmt: str) -> tuple[str, dict]:
        try:
            try:
                body = _text_file(path, "utf-8")
            except UnicodeDecodeError:
                body = _text_file(path, "latin-1")
        except Exception as error:
            return _read_error(error)
        return body, {"format": fmt}

    def _read_pdf(self, path: str) -> tuple[str, dict]:
        if pdfplumber is None:
            return _missing_reader("PDF", "pdfplumber")
        try:
            file_size = os.path.getsize(path)
            self._check_pdf_limit(
                "file bytes", file_size, self.pdf_limits.max_file_bytes
            )
            with pdfplumber.open(path) as document:
                page_count = len(document.pages)
                self._check_pdf_limit(
                    "page count", page_count, self.pdf_limits.max_pages
                )
                pages = []
                text_chars = 0
                for page in document.pages:
                    text = page.extract_text() or ""
                    text_chars += len(text)
                    self._check_pdf_limit(
                        "text characters", text_chars, self.pdf_limits.max_text_chars
                    )
                    pages.append(text)

                scanned = [
                    index for index, text in enumerate(pages) if not text.strip()
                ]
                metadata = {
                    "format": "pdf",
                    "page_count": page_count,
                    "scanned_page_count": len(scanned),
                }
                if not scanned:
                    return "\n".join(pages), metadata
                if self.ocr_provider is None:
                    metadata.update(ocr_required=True, ocr_available=False)
                    if not any(text.strip() for text in pages):
                        message = "scanned PDF requires an available OCR provider"
                        metadata.update(
                            format="error", error=message, error_kind="ocr_unavailable"
                        )
                        return f"Error reading PDF: {message}", metadata
                    return "\n".join(pages), metadata

                self._check_pdf_limit(
                    "raster page count",
                    len(scanned),
                    self.pdf_limits.max_raster_pages,
                )
                self._ocr_scanned_pages(document.pages, pages, scanned)
                metadata.update(
                    ocr_required=True,
                    ocr_available=True,
                    ocr_used=True,
                    ocr_page_count=len(scanned),
                )
            return "\n".join(pages), metadata
        except PdfResourceLimitError as error:
            return f"Error reading PDF: {error}", {
                "format": "error",
                "error": str(error),
                "error_kind": "resource_limit",
                "resource": error.resource,
                "actual": error.actual,
                "limit": error.limit,
            }
        except Exception as error:
            text = self._salvage_as_text(path)
            return (
                (text, {"format": "text", "recovered_from": "pdf"})
                if text is not None
                else _read_error(error)
            )

    def _ocr_scanned_pages(
        self, pdf_pages, pages: list[str], scanned: list[int]
    ) -> None:
        limits = self.pdf_limits
        provider = self.ocr_provider
        if provider is None:
            raise RuntimeError("OCR provider is unavailable")
        raster_pixels = 0
        raster_bytes = 0
        text_chars = sum(len(text) for text in pages)
        with tempfile.TemporaryDirectory(prefix="gideon-pdf-ocr-") as directory:
            for index in scanned:
                page = pdf_pages[index]
                width = max(1, round(float(page.width) * limits.raster_dpi / 72))
                height = max(1, round(float(page.height) * limits.raster_dpi / 72))
                page_pixels = width * height
                self._check_pdf_limit(
                    "page raster pixels", page_pixels, limits.max_page_pixels
                )
                raster_pixels += page_pixels
                self._check_pdf_limit(
                    "total raster pixels", raster_pixels, limits.max_raster_pixels
                )

                image_path = os.path.join(directory, f"page-{index + 1}.png")
                page.to_image(resolution=limits.raster_dpi).save(
                    image_path, format="PNG"
                )
                raster_bytes += os.path.getsize(image_path)
                self._check_pdf_limit(
                    "raster bytes", raster_bytes, limits.max_raster_bytes
                )
                text = provider.ocr(image_path, page_number=index + 1)
                if not isinstance(text, str):
                    raise TypeError("OCR provider must return text")
                text_chars += len(text)
                self._check_pdf_limit(
                    "text characters", text_chars, limits.max_text_chars
                )
                pages[index] = text.strip()

    @staticmethod
    def _check_pdf_limit(resource: str, actual: int, limit: int) -> None:
        if actual > limit:
            raise PdfResourceLimitError(resource, actual, limit)

    @staticmethod
    def _salvage_as_text(path: str) -> str | None:
        try:
            with open(path, "rb") as stream:
                raw = stream.read(200_000)
        except OSError:
            return None
        if not raw:
            return None
        try:
            text = raw.decode("utf-8")
        except UnicodeDecodeError:
            return None
        sample = text[:2000]
        usable = sum(
            character.isprintable() or character in "\n\r\t " for character in sample
        )
        return text if usable / len(sample) >= 0.9 else None

    def _read_pptx(self, path: str) -> tuple[str, dict]:
        if Presentation is None:
            return _missing_reader("PPTX", "python-pptx")
        try:
            presentation = Presentation(path)
            sections = []
            for index, slide in enumerate(presentation.slides, 1):
                title, body = "", []
                for shape in slide.shapes:
                    if not shape.has_text_frame:
                        continue
                    text = shape.text_frame.text.strip()
                    if shape == slide.shapes.title:
                        title = text
                    else:
                        body.append(text)
                section = f"## Slide {index}: {title}\n" + "\n".join(body)
                notes = (
                    slide.notes_slide.notes_text_frame
                    if slide.has_notes_slide
                    else None
                )
                if notes and notes.text.strip():
                    section += "\n" + notes.text.strip()
                sections.append(section)
            return "\n\n".join(sections), {
                "format": "pptx",
                "slide_count": len(presentation.slides),
            }
        except Exception as error:
            return _read_error(error)

    def _read_docx(self, path: str) -> tuple[str, dict]:
        if Document is None:
            return _missing_reader("DOCX", "python-docx")
        try:
            document = Document(path)
            parts = [_paragraph_markup(paragraph) for paragraph in document.paragraphs]
            for table in document.tables:
                rows = _render_docx_table(table)
                if rows:
                    parts.extend(["", *rows])
            return "\n".join(parts), dict(
                format="docx",
                content_type="markdown",
                paragraph_count=len(document.paragraphs),
                table_count=len(document.tables),
            )
        except Exception as error:
            return _read_error(error)

    def _read_xlsx(self, path: str) -> tuple[str, dict]:
        if _load_workbook is None:
            return _missing_reader("XLSX", "openpyxl")
        try:
            workbook = _load_workbook(path, read_only=True, data_only=True)
            try:
                sections, count = [], 0
                for sheet in workbook.worksheets:
                    rows = _nonempty_cells(sheet.iter_rows(values_only=True))
                    if rows:
                        count += len(rows)
                        sections.append(
                            "\n".join([f"## {sheet.title}", *_table_lines(rows)])
                        )
                metadata = dict(
                    format="xlsx",
                    content_type="markdown",
                    sheet_count=len(workbook.worksheets),
                    row_count=count,
                )
            finally:
                workbook.close()
            return "\n\n".join(sections), metadata
        except Exception as error:
            return _read_error(error, "spreadsheet")

    def _read_csv(self, path: str) -> tuple[str, dict]:
        import csv

        kind = "tsv" if Path(path).suffix.lower() == ".tsv" else "csv"
        try:
            try:
                stream = open(path, newline="", encoding="utf-8")
            except UnicodeDecodeError:
                stream = open(path, newline="", encoding="latin-1")
            with stream:
                rows = _nonempty_cells(
                    csv.reader(stream, delimiter="\t" if kind == "tsv" else ",")
                )
        except Exception as error:
            return _read_error(error, kind.upper())
        metadata = dict(format=kind, content_type="markdown", row_count=len(rows))
        if not rows:
            return "", metadata
        width = max(map(len, rows))
        shown = [
            row + [""] * (width - len(row)) for row in rows[: self._CSV_MAX_TABLE_ROWS]
        ]
        lines = _table_lines(shown, escape=True)
        if len(shown) < len(rows):
            lines.append(f"\n_…{len(rows) - len(shown)} more rows_")
        return "\n".join(lines), metadata

    def _read_html(self, path: str) -> tuple[str, dict]:
        try:
            html = _text_file(path, "utf-8")
        except UnicodeDecodeError:
            html = _text_file(path, "latin-1")
        except Exception as error:
            return _read_error(error)
        return html_to_prose(html), {"format": "html"}


def _nonempty_cells(values):
    result = []
    for row in values:
        cells = ["" if cell is None else str(cell) for cell in row]
        if any(cell.strip() for cell in cells):
            result.append(cells)
    return result


def _paragraph_markup(paragraph):
    style = paragraph.style.name if paragraph.style else ""
    if not style.startswith("Heading"):
        return paragraph.text
    try:
        level = int(style.split()[-1])
    except (ValueError, IndexError):
        level = 1
    return "#" * level + " " + paragraph.text


def html_to_prose(html: str) -> str:
    from gideon.cognition.knowledge.connectors.base import strip_html_chrome

    content = strip_html_chrome(html or "")
    if _html2text_mod is None:
        return re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", content)).strip()
    converter = _html2text_mod.HTML2Text()
    converter.ignore_links, converter.ignore_images = False, True
    return converter.handle(content)


@dataclass(frozen=True)
class PdfLine:
    page: int
    text: str
    size: float
    char_count: int


@dataclass(frozen=True)
class PdfStructure:
    pages: tuple[str, ...] = ()
    lines: tuple[PdfLine, ...] = ()
    outline: tuple[str, ...] = field(default=())


_LINE_TOLERANCE_PT = 1.5


def read_pdf_structure(path: str) -> PdfStructure | None:
    if pdfplumber is None:
        return None
    if is_sensitive_path(path):
        raise PermissionError(f"Refusing to read sensitive path: {path}")
    try:
        with pdfplumber.open(path) as pdf:
            pages, lines = [], []
            for page_index, page in enumerate(pdf.pages):
                pages.append(page.extract_text() or "")
                lines.extend(_lines_for_page(page, page_index))
            return PdfStructure(tuple(pages), tuple(lines), _outline_titles(pdf))
    except Exception:
        return None


def _lines_for_page(page, index: int) -> list[PdfLine]:
    chars = sorted(
        (char for char in (page.chars or []) if (char.get("text") or "") != ""),
        key=lambda char: (
            round(float(char.get("top") or 0.0), 2),
            float(char.get("x0") or 0.0),
        ),
    )
    groups, current, baseline = [], [], None
    for char in chars:
        top = float(char.get("top") or 0.0)
        if baseline is None or abs(top - baseline) > _LINE_TOLERANCE_PT:
            if current:
                groups.append(current)
            current, baseline = [], top
        current.append(char)
    if current:
        groups.append(current)
    return [line for group in groups if (line := _line_from(group, index)) is not None]


def _line_from(chars: list[dict], index: int) -> PdfLine | None:
    text = "".join(str(char.get("text") or "") for char in chars).strip()
    if not text:
        return None
    size = max(float(char.get("size") or 0.0) for char in chars)
    return PdfLine(index, text, round(size, 2), len(chars))


def _outline_titles(pdf) -> tuple[str, ...]:
    try:
        entries = pdf.doc.get_outlines()
    except Exception:
        return ()
    titles = []
    for entry in entries:
        try:
            value = entry[1]
        except (IndexError, TypeError):
            continue
        if isinstance(value, bytes):
            value = value.decode("utf-8", "replace")
        title = " ".join(str(value or "").split())
        if title:
            titles.append(title)
    return tuple(titles)
