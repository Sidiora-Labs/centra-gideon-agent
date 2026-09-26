"""DeckModel → .pptx bytes (python-pptx).

Uses the template's real TITLE + BODY placeholders rather than free-floating text boxes.
That matters for more than tidiness: the pptx reader identifies a slide's title via
``slide.shapes.title``, so a deck built from text boxes would round-trip with every title
lost — and so would PowerPoint's own outline view.

**A bullet's depth is written, not assumed.** This writer used to pin ``level = 0`` on
every appended paragraph, so a nested outline came out of the model flat: the depth was
not lost in the file, it was never written. ``documents/pptx_parser.py`` reads the same
attribute back, which is what makes depth a round trip rather than a one-way flatten.
"""

from __future__ import annotations

import io

from gideon.workspace.documents.model import DeckModel, ShapeBox, Slide
from gideon.workspace.documents.pptx_shapes import body_placeholder
from gideon.workspace.documents.registry import register_writer

_LAYOUT_TITLE = 0
_LAYOUT_TITLE_CONTENT = 1
_LAYOUT_TITLE_ONLY = 5
_TEMPLATE_KEY = "gideon-template="


def render_pptx(model: object) -> bytes:
    from pptx.util import Inches

    if not isinstance(model, DeckModel):
        raise TypeError("pptx writer expects a DeckModel")
    prs = _presentation_for(model)
    if model.width_in > 0:
        prs.slide_width = Inches(model.width_in)
    if model.height_in > 0:
        prs.slide_height = Inches(model.height_in)
    if model.title:
        layout = prs.slide_layouts[_LAYOUT_TITLE]
        slide = prs.slides.add_slide(layout)
        if slide.shapes.title is not None:
            slide.shapes.title.text = model.title
    for entry in model.slides:
        _add_slide(prs, entry)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _presentation_for(model: DeckModel):
    from pptx import Presentation

    if not model.template_slug:
        return Presentation()
    from gideon.workspace.artifacts import registry as artifact_registry

    prov = artifact_registry.get_provider()
    art = prov.get(model.template_slug) if prov is not None else None
    if art is None or art.kind != "pptx":
        raise ValueError(f"PPTX template artifact {model.template_slug!r} is unavailable")
    version = model.template_version or art.version
    result = prov.raw_bytes(art.slug, version=version)
    if result is None:
        raise ValueError(f"PPTX template artifact {art.slug!r} version {version} is unavailable")
    prs = Presentation(io.BytesIO(result[0]))
    for slide_id in list(prs.slides._sldIdLst):
        prs.part.drop_rel(slide_id.rId)
        prs.slides._sldIdLst.remove(slide_id)
    model.template_version = version
    keywords = [part.strip() for part in (prs.core_properties.keywords or "").split(";")
                if part.strip() and not part.strip().startswith(_TEMPLATE_KEY)]
    prs.core_properties.keywords = "; ".join([*keywords, f"{_TEMPLATE_KEY}{art.slug}@{version}"])
    return prs


def _layout_for(prs, entry: Slide):
    """The slide's declared layout, or the one its content asks for.

    Resolved from the REAL template by name rather than from `DECK_LAYOUTS` by index: a
    template whose layouts were renamed or reordered would otherwise silently produce a
    different slide than the one the model names.
    """
    if entry.layout:
        for layout in prs.slide_layouts:
            if layout.name == entry.layout:
                return layout
    return prs.slide_layouts[
        _LAYOUT_TITLE_CONTENT if entry.bullets else _LAYOUT_TITLE_ONLY
    ]


def _add_slide(prs, entry: Slide) -> None:
    slide = prs.slides.add_slide(_layout_for(prs, entry))
    if slide.shapes.title is not None:
        slide.shapes.title.text = entry.title or ""
        _place(slide.shapes.title, entry.title_box)

    if entry.bullets:
        body = body_placeholder(slide)
        if body is not None:
            frame = body.text_frame
            frame.text = entry.bullets[0].text
            frame.paragraphs[0].level = entry.bullets[0].level
            for bullet in entry.bullets[1:]:
                para = frame.add_paragraph()
                para.text = bullet.text
                para.level = bullet.level
            _place(body, entry.body_box)

    if entry.artifact_slug:
        _add_artifact_picture(prs, slide, entry)

    notes_text = entry.notes
    if entry.artifact_slug:
        notes_text = (notes_text + f"\n[image: {entry.artifact_slug}]").strip()
    if notes_text:
        slide.notes_slide.notes_text_frame.text = notes_text


def _add_artifact_picture(prs, slide, entry: Slide) -> None:
    from PIL import Image

    from gideon.workspace.artifacts import registry as artifact_registry

    prov = artifact_registry.get_provider()
    art = prov.get(entry.artifact_slug) if prov is not None else None
    if art is None or art.kind != "image":
        raise ValueError(f"image artifact {entry.artifact_slug!r} is unavailable")
    result = prov.raw_bytes(art.slug)
    if result is None:
        raise ValueError(f"image artifact {art.slug!r} has no readable bytes")
    with Image.open(io.BytesIO(result[0])) as image:
        width, height = image.size
        if width < 1 or height < 1 or width * height > 40_000_000:
            raise ValueError(f"image artifact {art.slug!r} has unsupported dimensions")
        picture_bytes = io.BytesIO()
        image.save(picture_bytes, format="PNG")
    picture_bytes.seek(0)
    slide_width, slide_height = int(prs.slide_width), int(prs.slide_height)
    box_width = int(slide_width * (0.42 if entry.bullets else 0.76))
    box_height = int(slide_height * 0.62)
    scale = min(box_width / width, box_height / height)
    picture_width, picture_height = int(width * scale), int(height * scale)
    left = int(slide_width * (0.54 if entry.bullets else 0.12)) + (box_width - picture_width) // 2
    top = int(slide_height * 0.27) + (box_height - picture_height) // 2
    picture = slide.shapes.add_picture(picture_bytes, left, top, width=picture_width, height=picture_height)
    picture.name = f"Gideon image artifact {art.slug}"
    if entry.bullets and not entry.body_box.placed:
        body = body_placeholder(slide)
        if body is not None:
            body.width = min(int(body.width), int(slide_width * 0.46))


def _place(shape, box: ShapeBox) -> None:
    """Pin *shape* to *box*, or leave it inheriting the layout's position.

    An unplaced box writes NOTHING: a placeholder with no explicit geometry inherits from
    its layout, and pinning the inherited value would turn "wherever the layout puts it"
    into a frozen position on the first save.
    """
    from pptx.util import Inches

    if not box.placed:
        return
    shape.left = Inches(box.left_in)
    shape.top = Inches(box.top_in)
    shape.width = Inches(box.width_in)
    shape.height = Inches(box.height_in)


register_writer("pptx", render_pptx)
